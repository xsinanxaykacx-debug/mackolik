import os
from pathlib import Path
import duckdb
from openpyxl import load_workbook

print("DuckDB surumu:", duckdb.__version__)

roots = [
    Path.home() / "Downloads",
    Path.home() / "Desktop",
    Path.home() / "Documents",
    Path.home() / "OneDrive",
]

files = []
seen = set()

for root in roots:
    if not root.exists():
        continue
    try:
        for p in root.rglob("DATA.xlsx"):
            if p.is_file() and str(p).lower() not in seen:
                files.append(p)
                seen.add(str(p).lower())
    except (PermissionError, OSError):
        pass

if not files:
    print("\nDATA.xlsx bulunamadi.")
    print("Dosyanin tam yolunu kontrol et.")
    raise SystemExit(1)

print("\nBulunan dosyalar:")
for i, p in enumerate(files, 1):
    print(f"{i}. {p}")

if len(files) > 1:
    print("\nBirden fazla dosya bulundu. Dogru dosyayi secmeden aktarim yapmiyoruz.")
    raise SystemExit(0)

path = files[0]
print("\nSecilen dosya:", path)
print("Dosya boyutu (MB):", round(path.stat().st_size / (1024 * 1024), 2))

wb = load_workbook(
    filename=str(path),
    read_only=True,
    data_only=True
)

print("\n=== SAYFA BILGILERI ===")

for ws in wb.worksheets:
    print(f"\nSayfa: {ws.title}")
    print("Excel boyutu:", ws.max_row, "satir x", ws.max_column, "sutun")

    first_rows = ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row or 1, 4),
        values_only=True
    )

    for row_num, row in enumerate(first_rows, 1):
        values = list(row)
        while values and values[-1] is None:
            values.pop()
        print(f"Satir {row_num}:", values)

wb.close()

print("\n=== ON INCELEME TAMAMLANDI ===")
print("Excel dosyasi degistirilmedi.")